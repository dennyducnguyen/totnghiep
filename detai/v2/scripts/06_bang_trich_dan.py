# -*- coding: utf-8 -*-
"""
Tạo bảng đối chiếu trích dẫn:  v2/Bang-doi-chieu-trich-dan.xlsx
Mỗi dòng = một tài liệu trong danh mục Tài liệu tham khảo của bài báo (scripts/noi_dung.js), gồm:
  tên bài, tác giả (năm), câu hỏi nghiên cứu và mô hình của bài gốc, đoạn nội dung gốc cần dùng,
  các câu trong bài báo có trích dẫn tài liệu đó, dạng trích dẫn APA trong bài, mục xuất hiện, APA đầy đủ,
  nguồn đối chiếu / trạng thái kiểm chứng.
Nguồn nội dung gốc: doi_chieu_citations.xlsx (25 tài liệu đã đối chiếu nguyên văn),
  Nghien_Cuu_Tai_Lieu_EBI.xlsx (sheet Phieu_Doc_12_Tai_Lieu, Tai_Lieu_Thay_The_2026), phiếu đọc 3 bài mẫu,
  và phần tóm tắt của tác giả cho các tài liệu kinh điển còn lại (ghi rõ trạng thái).
"""
import os, re, json
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))          # v2/

def tim_file(ten):
    """Tìm file dữ liệu gốc: trong thư mục phiên bản này, nếu không có thì ở thư mục cha (detai/)."""
    for d in (ROOT, os.path.dirname(ROOT)):
        p = os.path.join(d, ten)
        if os.path.exists(p): return p
    raise FileNotFoundError(ten)

ND = open(os.path.join(ROOT, "scripts", "noi_dung.js"), encoding="utf-8").read()

# ----------------------------------------------------------------------------------------
# 1. Đọc bài báo: danh mục TLTK và các đoạn văn kèm mục
# ----------------------------------------------------------------------------------------
refs_raw = re.search(r't: "refs", items: \[(.*?)\n  \] \}', ND, re.S).group(1)
REFS = [r.replace("\\'", "'") for r in re.findall(r'^\s+"(.*)",\s*$', refs_raw, re.M)]
blocks = []
muc = ""
for m in re.finditer(r'\{ t: "(h1|h2|p|abstract|bullets|fig)", (?:text|items|file): (.*?) \},?\n', ND, re.S):
    kind, body = m.group(1), m.group(2)
    if kind in ("h1", "h2"):
        muc = re.match(r'"(.*?)"', body).group(1); continue
    if kind == "fig":
        cap = re.search(r'caption: "(.*?)"', body).group(1); blocks.append((muc, cap)); continue
    texts = re.findall(r'"((?:[^"\\]|\\.)*)"', body)
    for t in texts: blocks.append((muc or "Tóm tắt", t))
def clean(t): return re.sub(r"\*\*?|\^\{|_\{|\}", "", t)
SENTS = []
for muc, t in blocks:
    t = clean(t)
    for s in re.split(r"(?<=[.!?])\s+(?=[A-ZĐÂÊÔƠƯÁÀẢÃẠĂẤẦẨẪẬẮẰẲẴẶÉÈẺẼẸẾỀỂỄỆÍÌỈĨỊÓÒỎÕỌỐỒỔỖỘỚỜỞỠỢÚÙỦŨỤỨỪỬỮỰÝỲỶỸỴ(*])", t):
        if s.strip(): SENTS.append((muc, s.strip()))

# ----------------------------------------------------------------------------------------
# 2. Nguồn đối chiếu bên ngoài
# ----------------------------------------------------------------------------------------
wb_dc = load_workbook(tim_file("doi_chieu_citations.xlsx"), read_only=True, data_only=True)
DC = {}
for row in wb_dc.worksheets[0].iter_rows(min_row=2, values_only=True):
    if row[0] is not None: DC[int(row[0])] = dict(goc=row[4], trich=row[5], apa=row[8], danh_gia=row[6])
wb_nc = load_workbook(tim_file("Nghien_Cuu_Tai_Lieu_EBI.xlsx"), read_only=True, data_only=True)
PD = {}
for name in ["Phieu_Doc_12_Tai_Lieu", "Tai_Lieu_Thay_The_2026"]:
    for row in wb_nc[name].iter_rows(min_row=4, values_only=True):
        if row[0] and re.match(r"^[A-Z]\d[a-z]?$", str(row[0])):
            PD[row[0]] = dict(rq=row[3], rq_en=row[4], mh=row[7], pp=row[12], kq=row[14], kf_en=row[15])

# ----------------------------------------------------------------------------------------
# 3. Bảng tra: mỗi tài liệu -> khóa tìm trong bài, nguồn đối chiếu, nội dung bổ sung thủ công
#    key: chuỗi (regex) nhận diện trích dẫn trong bài; dc: STT trong doi_chieu; pd: mã phiếu đọc
#    goc/rq/mh: nội dung thủ công (dùng khi không có nguồn ở trên); tt: trạng thái
# ----------------------------------------------------------------------------------------
TT_NGUYEN_VAN = "Nguyên văn đã đối chiếu (doi_chieu_citations.xlsx)"
TT_PHIEU = "Từ phiếu đọc (Nghien_Cuu_Tai_Lieu_EBI.xlsx) – trích tiếng Anh từ bài/abstract, nên đối chiếu lại số trang"
TT_MAU = "Đã đọc toàn văn PDF trong bai-bao-mau (trích nguyên văn)"
TT_TOM_TAT = "Tóm tắt của tác giả từ nội dung gốc – CHƯA trích nguyên văn, cần mở tài liệu để lấy câu gốc"

M = [
 dict(key=r"Abd Rashid", dc=17, rq="Khoảng cách số giữa và trong năm nước ASEAN (Việt Nam, Philippines, Lào, Campuchia, Indonesia) biểu hiện ra sao và chính sách nào thu hẹp được?", mh="Phân tích so sánh chính sách và dữ liệu thứ cấp đa quốc gia; không có mô hình kinh tế lượng.", tt=TT_NGUYEN_VAN),
 dict(key=r"Anselin", pd="D1", tt=TT_PHIEU),
 dict(key=r"Aumeboonsuke", dc=13, rq="Các nền kinh tế EU và ASEAN đang hội tụ hay phân kỳ về phát triển số?", mh="Kiểm định hội tụ beta và sigma trên chỉ số phát triển số cấp quốc gia, so sánh trước/sau COVID-19.", tt=TT_NGUYEN_VAN),
 dict(key=r"Barro", dc=7, pd="B1", tt=TT_NGUYEN_VAN),
 dict(key=r"Billon", pd="C2", tt=TT_PHIEU),
 dict(key=r"Blinder", pd="B4", tt=TT_PHIEU),
 dict(key=r"Bộ Công Thương", dc=1, rq="(Văn bản chính sách) Mục tiêu, nhiệm vụ phát triển TMĐT giai đoạn 2026–2030.", mh="Không áp dụng (văn bản pháp quy).", tt=TT_NGUYEN_VAN),
 dict(key=r"Bruno", dc=14, rq="Chỉ số DESI có dư thừa chỉ tiêu không, và có thể xây dựng bản rút gọn dùng được ở cấp vùng (20 vùng của Ý) hay không?", mh="Quy trình 4 bước: tương quan cặp chỉ tiêu → PCA theo nhóm → chọn chỉ tiêu đại diện, phân bổ lại trọng số → xác minh bằng tương quan và Normalized Spearman's Footrule; ứng dụng cho vùng Ý; khung chính sách μ×σ.",
      goc='"Such an approach seems to respond to the assumption that a higher number of indicators should result in more accurate CIs." (tr. 1); "The present study aims at giving a twofold contribution." (tr. 2); kết quả: DESI rút gọn 15/37 chỉ tiêu, tương quan điểm tổng 0,95, Lombardia 48,66% – Molise 31,03%.', tt=TT_MAU),
 dict(key=r"Caragliu", dc=25, rq="Chất lượng thể chế vùng ảnh hưởng thế nào đến khoảng cách số đô thị ở châu Âu?", mh="Hồi quy dữ liệu 181 thành phố châu Âu với chỉ số chất lượng thể chế vùng (EQI) làm biến giải thích.", tt=TT_NGUYEN_VAN),
 dict(key=r"Chinn", pd="C1", tt=TT_PHIEU),
 dict(key=r"Cowell", pd="B3", tt=TT_PHIEU),
 dict(key=r"Davis \(1989\)|Davis, 1989", rq="Điều gì khiến người dùng chấp nhận hay từ chối một công nghệ thông tin?", mh="Mô hình chấp nhận công nghệ (TAM): nhận thức hữu ích và nhận thức dễ sử dụng → thái độ → ý định → hành vi sử dụng; kiểm định bằng thang đo và hồi quy trên hai nghiên cứu thực nghiệm.",
      goc="Nội dung dùng trong bài: hai cấu trúc 'perceived usefulness' và 'perceived ease of use' là yếu tố quyết định cơ bản của việc chấp nhận công nghệ ở cấp người dùng/doanh nghiệp.", tt=TT_TOM_TAT),
 dict(key=r"Fan và cộng sự", pd="R3a", tt=TT_PHIEU),
 dict(key=r"Galperin", pd="R5c", tt=TT_PHIEU),
 dict(key=r"Hao và cộng sự", pd="R4b", tt=TT_PHIEU),
 dict(key=r"VECOM", dc=3, rq="(Báo cáo dữ liệu) Xếp hạng mức độ phát triển TMĐT của các tỉnh, thành phố hằng năm.", mh="Chỉ số tổng hợp có trọng số từ các trụ cột NNL&HT, B2C, B2B (và G2B đến 2020), dựa trên khảo sát doanh nghiệp và dữ liệu thứ cấp.", tt=TT_NGUYEN_VAN),
 dict(key=r"ITU", dc=5, rq="(Báo cáo thống kê) Hiện trạng phát triển số toàn cầu năm 2025.", mh="Thống kê mô tả từ dữ liệu ITU.", tt=TT_NGUYEN_VAN),
 dict(key=r"Jann", rq="Làm thế nào ước lượng và suy diễn thống kê cho phân tách Blinder–Oaxaca một cách nhất quán, kể cả việc chọn hệ số tham chiếu?", mh="Trình bày phân tách hai và ba thành phần, các lựa chọn véc-tơ hệ số tham chiếu (trong đó có hệ số từ mô hình gộp có biến chỉ báo nhóm), sai số chuẩn; lệnh oaxaca trong Stata.",
      goc="Nội dung dùng trong bài: khi dùng hệ số của mô hình gộp làm tham chiếu, cần đưa biến chỉ báo nhóm vào mô hình gộp để tránh chuyển một phần khác biệt nhóm vào phần nguồn lực.", tt=TT_TOM_TAT),
 dict(key=r"Laskar", dc=15, rq="Xã hội số ở Ấn Độ hình thành ra sao và khoảng cách số nông thôn–thành thị khác nhau thế nào?", mh="So sánh mô tả nông thôn – thành thị từ dữ liệu khảo sát quốc gia.", tt=TT_NGUYEN_VAN),
 dict(key=r"Leogrande", pd="R4a", tt=TT_PHIEU),
 dict(key=r"Lerman", dc=8, rq="Đóng góp của từng nguồn thu nhập vào bất bình đẳng tổng thể (Gini) là bao nhiêu?", mh="Phân rã Gini theo nguồn: G = Σ S_k G_k R_k (tỷ trọng × Gini nguồn × tương quan Gini); áp dụng cho dữ liệu Mỹ.", tt=TT_NGUYEN_VAN),
 dict(key=r"Li và cộng sự \(2025\)", pd="R2b", tt=TT_PHIEU),
 dict(key=r"Lin và cộng sự", pd="R5b", tt=TT_PHIEU),
 dict(key=r"Norris", dc=9, pd="A2", tt=TT_NGUYEN_VAN),
 dict(key=r"Oaxaca \(1973\)|Oaxaca, 1973", pd="B4", tt=TT_PHIEU),
 dict(key=r"Özcan Alp", dc=16, rq="Khoảng cách số giữa các vùng/tỉnh của Thổ Nhĩ Kỳ lớn đến đâu?", mh="Chỉ số tổng hợp khoảng cách số cấp tỉnh và phân tích so sánh vùng.", tt=TT_NGUYEN_VAN),
 dict(key=r"Phillips", pd="D1", tt=TT_PHIEU),
 dict(key=r"Qiu và cộng sự", dc=18, rq="Khoảng cách số có làm rộng khoảng cách thu nhập giữa các vùng của Trung Quốc không?", mh="Chỉ số Theil đo bất bình đẳng ở 280 thành phố; hồi quy dữ liệu bảng khoảng cách thu nhập lên khoảng cách số.", tt=TT_NGUYEN_VAN),
 dict(key=r"Raihan", dc=19, rq="Các chiều và rào cản của bất bình đẳng số là gì?", mh="Tổng quan tích hợp có hệ thống (systematic integrative review).", tt=TT_NGUYEN_VAN),
 dict(key=r"Rogers", pd="A4", tt=TT_PHIEU),
 dict(key=r"Ruan và cộng sự", dc=20, rq="Phát triển kinh tế số tác động thế nào đến khoảng cách thu nhập vùng?", mh="Phân tách Gini đa chiều (xấp xỉ Kmenta) và mô hình ngưỡng.", tt=TT_NGUYEN_VAN),
 dict(key=r"Sala-i-Martin \(1996\)|Sala-i-Martin, 1996", pd="B2", tt=TT_PHIEU),
 dict(key=r"Theil \(1967\)|Theil, 1967", pd="B3", tt=TT_PHIEU),
 dict(key=r"Thủ tướng Chính phủ", dc=2, rq="(Văn bản chính sách) Mục tiêu phát triển TMĐT quốc gia 2021–2025.", mh="Không áp dụng (văn bản pháp quy).", tt=TT_NGUYEN_VAN),
 dict(key=r"Tornatzky", rq="Những yếu tố nào ở cấp tổ chức quyết định việc áp dụng đổi mới công nghệ?", mh="Khung Công nghệ – Tổ chức – Môi trường (TOE): ba bối cảnh ảnh hưởng đến quyết định ứng dụng công nghệ của doanh nghiệp.",
      goc="Nội dung dùng trong bài: quyết định ứng dụng công nghệ của doanh nghiệp phụ thuộc đồng thời vào bối cảnh công nghệ, tổ chức và môi trường (cạnh tranh, khách hàng, quy định).", tt=TT_TOM_TAT),
 dict(key=r"UNCTAD", dc=24, rq="(Báo cáo) Xu hướng kinh tế số toàn cầu 2024.", mh="Thống kê mô tả và tổng hợp chính sách.", tt=TT_NGUYEN_VAN),
 dict(key=r"van Deursen & Helsper|van Deursen và Helsper", nam="2015", dc=11, rq="Ai hưởng lợi nhiều nhất từ việc trực tuyến (khoảng cách số cấp ba)?", mh="Khảo sát người dùng Internet Hà Lan; hồi quy lợi ích thu được lên đặc điểm kinh tế – xã hội và kỹ năng.", tt=TT_NGUYEN_VAN),
 dict(key=r"van Deursen & van Dijk|van Deursen và van Dijk", nam="2014", dc=10, rq="Khi tiếp cận vật lý gần phổ cập, khoảng cách số có chuyển sang khác biệt về cách sử dụng không?", mh="Khảo sát đại diện dân số Hà Lan; phân tích khác biệt sử dụng theo học vấn, tuổi, giới.", tt=TT_NGUYEN_VAN),
 dict(key=r"van Dijk \(2006|van Dijk, 2006|van Dijk, 2006, 2020|van Dijk \(2006, 2020", pd="A1", tt=TT_PHIEU),
 dict(key=r"van Dijk \(2020|van Dijk, 2020|van Dijk, 2006, 2020|van Dijk \(2006, 2020", dc=12, rq="Khoảng cách số biến đổi thế nào sau 25 năm và tầng kết quả (third-level) biểu hiện ra sao?", mh="Tổng hợp lý thuyết: khung tiếp cận đa tầng (động cơ – vật chất – kỹ năng – sử dụng – kết quả) và lý thuyết nguồn lực – chiếm dụng (resources and appropriation theory).", tt=TT_NGUYEN_VAN),
 dict(key=r"Vicente", pd="C3", tt=TT_PHIEU),
 dict(key=r"CIEM", dc=4, rq="(Báo cáo chính sách) Thực trạng và khuyến nghị phát triển kinh tế số Việt Nam.", mh="Phân tích chính sách định tính, dữ liệu thứ cấp.", tt=TT_NGUYEN_VAN),
 dict(key=r"Vu \(2026\)|Vu, 2026", pd="R6a", tt=TT_PHIEU),
 dict(key=r"Vu và cộng sự \(2025\)|Vu và cộng sự, 2025", pd="R6b", tt=TT_PHIEU),
 dict(key=r"Wang và cộng sự \(2021\)|Wang và cộng sự, 2021|Wang et al\., 2021", dc=22, rq="Ở cấp huyện Trung Quốc, kinh doanh trực tuyến (OBI) và mua sắm trực tuyến (OSI) lan tỏa theo giả thuyết khuếch tán đổi mới hay giả thuyết hiệu quả, và vì sao khác nhau?", mh="Bản đồ phân cấp → điểm nóng Getis-Ord Gi* → bản đồ hiệu số OSI−OBI → đường suy giảm theo khoảng cách từ lõi; OLS hai mô hình (N = 1.919 huyện, dữ liệu AliResearch 2015).",
      goc='"Our study suggests that online business trends are in line only with the innovation diffusion hypothesis, with marginal counties having lower levels of online business." (Abstract); "Affordability is, therefore, the main reason why some marginal areas have achieved a \'marginal revolution\' in terms of online shopping." (tr. 12)', tt=TT_MAU),
 dict(key=r"Wang \(2024\)|Wang, 2024", pd="R3b", tt=TT_PHIEU),
 dict(key=r"Warschauer", rq="Vì sao cách hiểu khoảng cách số như 'có/không có máy tính và Internet' là sai lệch, và hòa nhập xã hội qua công nghệ đòi hỏi những nguồn lực nào?", mh="Nghiên cứu định tính nhiều trường hợp (Ai Cập, Ấn Độ, Brazil, Trung Quốc, Mỹ); khung bốn nguồn lực: vật chất, số, nhân lực, xã hội.",
      goc="Nội dung dùng trong bài: khoảng cách số không phải là vấn đề nhị phân về thiết bị mà là vấn đề hòa nhập xã hội, phụ thuộc vào nguồn lực vật chất, số, nhân lực và xã hội.", tt=TT_TOM_TAT),
 dict(key=r"Williamson", dc=6, pd="A3", tt=TT_NGUYEN_VAN),
 dict(key=r"Wu và cộng sự", pd="R2a", tt=TT_PHIEU),
 dict(key=r"Zhao và cộng sự", pd="R5a", tt=TT_PHIEU),
 dict(key=r"Zhou và cộng sự", dc=23, rq="Khoảng cách số (DD) của Trung Quốc nên định nghĩa và đo thế nào; DD ở cấp quốc gia – vùng – tỉnh biến động ra sao 2010–2020; yếu tố nào quyết định DD cấp tỉnh?", mh="Chỉ số kinh tế số (entropy, 5 chiều × 33 chỉ tiêu, 298 thành phố); DD = Theil/MLD/Gini và phân rã nội/liên vùng; hồi quy hiệu ứng cố định hai chiều (27 tỉnh × 11 năm).",
      goc='"This study defines DD as the spatial imbalance of the digital economy." (tr. 3); "The degree of the national DD shows a fluctuating downward trend, from 0.4731 in 2010 to 0.2560 in 2020, which is a decrease rate of 45.89%." (tr. 17); tỷ lệ đóng góp nội vùng tăng từ ~0,61 lên ~0,76 (Fig. 4).', tt=TT_MAU),
]

# ----------------------------------------------------------------------------------------
# 4. Ghép: mỗi TLTK ↔ một dòng M (theo họ tác giả đầu + năm)
# ----------------------------------------------------------------------------------------
def tim_M(ref):
    """Chọn dòng M cho một TLTK: họ (từ khóa) phải xuất hiện trong PHẦN TÁC GIẢ của TLTK (không xét tiêu đề);
    nếu khóa có năm thì năm phải khớp; nhiều ứng viên thì lấy ứng viên khớp gần đầu chuỗi nhất (tác giả thứ nhất)."""
    nam = re.search(r"\((\d{4})", ref).group(1)
    tac_gia = ref[:re.search(r"\(\d{4}", ref).start()]          # phần tác giả (trước năm), gồm cả tên viết tắt như (VECOM)
    ung = []
    for m in M:
        ho = re.split(r"\\\(|\(|,| và | & | et al", m["key"])[0].strip()
        hit = re.search(r"(?<![\w])" + re.escape(ho) + r"(?![\w])", tac_gia, re.I)
        if not hit: continue
        years = re.findall(r"\d{4}", m["key"]) + ([m["nam"]] if m.get("nam") else [])
        if years and nam not in years: continue
        ung.append((hit.start(), m))
    return sorted(ung, key=lambda x: x[0])[0][1] if ung else None

def cau_trich(key, nam):
    pat = re.compile(key)
    out = []
    for muc, s in SENTS:
        if pat.search(s) and (nam in s or "cộng sự" in s or key.startswith("van Dijk")):
            if key.startswith("van Dijk") and not re.search(key, s): continue
            out.append((muc, s))
    return out

def dang_apa(key, sents):
    """Liệt kê các dạng trích dẫn in-text thực sự chứa tên tài liệu: (Tác giả, năm) hoặc Tác giả (năm)."""
    forms = set()
    for _, s in sents:
        for m in re.finditer(r"\(([^()]*?\d{4}[^()]*)\)", s):               # dạng trong ngoặc
            if re.search(key, m.group(0)): forms.add(m.group(0))
        for m in re.finditer(r"((?:[A-ZĐÖÜ][\w'’\-]*(?: [\w'’\-]+){0,5}?)) \((\d{4}[a-z]?(?:[,–] ?\d{4})*)\)", s):   # dạng tường thuật
            k = re.search(key, m.group(0))
            if k: forms.add(m.group(0)[k.start():])          # cắt bỏ chữ đứng trước tên tác giả
    return "; ".join(sorted(forms))

rows = []
for i, ref in enumerate(REFS, 1):
    m = tim_M(ref) or {}
    nam = re.search(r"\((\d{4})", ref).group(1)
    key = m.get("key", re.match(r"([^,(]+)", ref).group(1).strip())
    sents = cau_trich(key, nam)
    ten = re.sub(r"^.*?\(\d{4}[^)]*\)\.\s*", "", ref); ten = re.sub(r"\*", "", ten)
    tacgia = re.match(r"(.*?)\s*\(\d{4}", ref).group(1)
    dc = DC.get(m.get("dc")) if m.get("dc") else None
    pd_ = PD.get(m.get("pd")) if m.get("pd") else None
    rq = m.get("rq") or (pd_ and pd_["rq"]) or ""
    mh = m.get("mh") or (pd_ and f"{pd_['mh']}\n— Phương pháp: {pd_['pp']}") or ""
    goc = m.get("goc") or (dc and dc["goc"]) or (pd_ and f"{pd_['kf_en']}\n[Câu hỏi gốc EN: {pd_['rq_en']}]") or ""
    if not goc: goc = "(Chưa có đoạn gốc – cần bổ sung khi mở tài liệu)"
    tt = m.get("tt", TT_TOM_TAT)
    if dc and m.get("goc") is None: tt = TT_NGUYEN_VAN + (f" – đánh giá: {dc['danh_gia']}" if dc.get("danh_gia") else "")
    rows.append([i, ten.strip(), tacgia.strip(), nam, rq, mh, goc,
                 "\n\n".join(f"[{mu}] {s}" for mu, s in sents) or "(không tìm thấy câu trích dẫn – kiểm tra lại)",
                 dang_apa(key, sents), "; ".join(sorted({mu for mu, _ in sents})), re.sub(r"\*", "", ref), tt])

# ----------------------------------------------------------------------------------------
# 5. Ghi Excel
# ----------------------------------------------------------------------------------------
wb = Workbook(); ws = wb.active; ws.title = "Doi_chieu_trich_dan"
head = ["STT", "Tên bài báo / tài liệu", "Tác giả", "Năm", "Câu hỏi nghiên cứu của bài gốc", "Mô hình / phương pháp nghiên cứu của bài gốc",
        "Đoạn nội dung gốc cần sử dụng", "Nội dung trích dẫn trong bài (các câu có trích dẫn, kèm mục)", "Dạng trích dẫn APA trong bài",
        "Mục xuất hiện trong bài", "Tài liệu tham khảo APA đầy đủ", "Nguồn đối chiếu / trạng thái kiểm chứng"]
ws["A1"] = "BẢNG ĐỐI CHIẾU TRÍCH DẪN – bài báo “Khoảng cách thương mại điện tử giữa các tỉnh, thành phố Việt Nam giai đoạn 2017–2024” (phiên bản v2)"
ws["A1"].font = Font(bold=True, size=13)
ws["A2"] = ("Mỗi dòng là một tài liệu trong danh mục Tài liệu tham khảo. Cột 'Nội dung trích dẫn trong bài' được trích tự động từ scripts/noi_dung.js "
            "(mọi câu có tên tác giả và năm); cột 'Đoạn nội dung gốc' lấy từ doi_chieu_citations.xlsx (nguyên văn đã đối chiếu), phiếu đọc trong "
            "Nghien_Cuu_Tai_Lieu_EBI.xlsx, hoặc phiếu đọc 3 bài mẫu; dòng nào ghi 'Tóm tắt của tác giả' là chưa có câu nguyên văn, cần mở tài liệu để bổ sung.")
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top"); ws.merge_cells("A2:L2"); ws.row_dimensions[2].height = 48
fill = PatternFill("solid", fgColor="D9E2F3"); thin = Side(style="thin", color="BFBFBF")
for j, h in enumerate(head, 1):
    c = ws.cell(row=4, column=j, value=h); c.font = Font(bold=True); c.fill = fill
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center"); c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
mau_tt = {TT_NGUYEN_VAN: "C6EFCE", TT_MAU: "C6EFCE", TT_PHIEU: "FFEB9C", TT_TOM_TAT: "F8CBAD"}
for r, row in enumerate(rows, 5):
    for j, v in enumerate(row, 1):
        c = ws.cell(row=r, column=j, value=v); c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    tt = row[-1]; fill_tt = next((v for k, v in mau_tt.items() if tt.startswith(k)), None)
    if fill_tt: ws.cell(row=r, column=12).fill = PatternFill("solid", fgColor=fill_tt)
    ws.row_dimensions[r].height = 220
for j, w in enumerate([5, 38, 22, 6, 40, 42, 55, 60, 28, 22, 45, 34], 1): ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = "C5"; ws.auto_filter.ref = f"A4:L{4 + len(rows)}"

ws2 = wb.create_sheet("Chu_thich")
for i, t in enumerate([
    "Ý nghĩa màu ở cột 'Nguồn đối chiếu / trạng thái kiểm chứng':",
    "  Xanh lá: đoạn gốc là nguyên văn đã đối chiếu (doi_chieu_citations.xlsx) hoặc trích trực tiếp từ PDF trong bai-bao-mau.",
    "  Vàng: lấy từ phiếu đọc trong Nghien_Cuu_Tai_Lieu_EBI.xlsx (câu tiếng Anh từ bài/abstract); nên mở bài gốc để xác nhận số trang.",
    "  Cam: mới có tóm tắt của tác giả về nội dung gốc, chưa có câu nguyên văn – cần bổ sung.",
    "Cột 'Dạng trích dẫn APA trong bài' liệt kê các dạng in-text đã dùng, ví dụ (van Dijk, 2006) hoặc van Dijk (2006).",
    "Cột 'Mục xuất hiện' cho biết tài liệu được trích ở mục nào của bài (Tóm tắt, 1. Giới thiệu, 2.x, 5. Thảo luận...).",
    "Cách tạo lại file này: python scripts/06_bang_trich_dan.py (đọc scripts/noi_dung.js, doi_chieu_citations.xlsx, Nghien_Cuu_Tai_Lieu_EBI.xlsx).",
], 1):
    ws2.cell(row=i, column=1, value=t)
ws2.column_dimensions["A"].width = 140
out = os.path.join(ROOT, "Bang-doi-chieu-trich-dan.xlsx"); wb.save(out)
thieu = [r[2] for r in rows if r[7].startswith("(không")]
print(f"Đã ghi {out}: {len(rows)} tài liệu ({sum(1 for r in rows if r[4])} có RQ); không tìm thấy câu trích dẫn cho: {thieu}")
print("Trạng thái:", {k: sum(1 for r in rows if r[-1].startswith(k)) for k in mau_tt})
